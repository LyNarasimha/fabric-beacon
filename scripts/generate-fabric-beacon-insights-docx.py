#!/usr/bin/env python3
"""Build Fabric Beacon project insights Word doc (Paginate Forge–style stakeholder guide).

Requires: pip install python-docx matplotlib

Outputs:
  docs/Fabric-Beacon-Project-Insights.docx
  %USERPROFILE%/Downloads/Fabric-Beacon-Project-Insights.docx
"""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCS = ROOT / "docs" / "Fabric-Beacon-Project-Insights.docx"
OUT_DOWNLOADS = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Downloads" / "Fabric-Beacon-Project-Insights.docx"
ASSETS = OUT_DOWNLOADS.parent / "Fabric-Beacon-Guide-assets"

BRAND = "Fabric Beacon"
ACCENT = RGBColor(0x00, 0x78, 0xD4)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = ACCENT if level <= 2 else RGBColor(0x33, 0x33, 0x33)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "E8F4FC") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, label in enumerate(headers):
        hdr[i].text = label
        set_cell_shading(hdr[i], header_fill)
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
    doc.add_paragraph()


def draw_user_journey(path: Path) -> None:
    import matplotlib.patches as mpatches
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")

    def box(x, y, w, h, text, fc="#E8F4FC"):
        rect = mpatches.FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            linewidth=1.4,
            edgecolor="#0078D4",
            facecolor=fc,
        )
        ax.add_patch(rect)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, wrap=True)

    def arrow(x1, y1, x2, y2):
        ax.annotate(
            "",
            xy=(x2, y2),
            xytext=(x1, y1),
            arrowprops=dict(arrowstyle="->", color="#555555", lw=1.2),
        )

    box(0.3, 2.0, 2.0, 1.2, "Start app\nstart-dev.ps1\n(Demo or .env)", "#D4EDDA")
    box(2.6, 2.0, 2.0, 1.2, "Setup\nConnect + pick\n[QA] workspaces", "#E8F4FC")
    box(4.9, 2.0, 2.0, 1.2, "Sync catalog\nReports + models", "#E8F4FC")
    box(7.2, 2.0, 2.5, 1.2, "Jira Analyzer\nTicket → QA scope\nmatches only", "#FFF3CD")
    box(3.2, 0.35, 3.6, 1.0, "Discovery + Documentation\nBrowse, search, business docs", "#F5F5F5")

    arrow(2.3, 2.6, 2.6, 2.6)
    arrow(4.6, 2.6, 4.9, 2.6)
    arrow(6.9, 2.6, 7.2, 2.6)
    arrow(8.45, 2.0, 5.0, 1.35)

    ax.text(5, 4.5, f"{BRAND} — new user journey (web app)", ha="center", fontsize=12, fontweight="bold")
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_docx(journey_png: Path | None) -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(BRAND)
    tr.bold = True
    tr.font.size = Pt(26)
    tr.font.color.rgb = ACCENT
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Project insights — catalog discovery, documentation, Jira-scoped analysis")
    sr.font.size = Pt(13)
    sr.font.color.rgb = MUTED
    doc.add_paragraph(f"Version: {date.today().isoformat()}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Repo slug: pbi-discovery-platform · Public name: Fabric Beacon").alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )
    doc.add_page_break()

    add_heading(doc, "Who this document is for", 1)
    doc.add_paragraph(
        "Leadership, enablement, and report developers who need a single Word companion to the repo "
        "setup guide (docs/DOCUMENTATION.md). It summarizes what Fabric Beacon does, how teams onboard, "
        "what is already signed off in engineering, and what still requires your tenant (QA sync, staging cert)."
    )

    add_heading(doc, "Executive summary", 1)
    doc.add_paragraph(
        "Fabric Beacon is the self-service web platform for Microsoft Fabric and Power BI catalog work: "
        "scoped workspace discovery, semantic model inspection, business documentation, governance signals, "
        "and Jira ticket analysis that matches assets only inside the ticket’s resolved [QA] workspace — "
        "never a whole-tenant fallback."
    )
    doc.add_paragraph(
        "Engineering sign-off is GO — 100% for the defined new-user scope (About → Setup → Jira → "
        "Discovery/Documentation). Production rollout in your organization still needs live credentials, "
        "[QA] workspace sync, and optional staging certification per PLATFORM-NEXT-ACTIONS.md."
    )

    add_heading(doc, "Sign-off snapshot", 1)
    add_table(
        doc,
        ["Gate", "Status", "Evidence"],
        [
            ("New-user journey (§0 DOCUMENTATION.md)", "Shipped", "About, Setup, Jira, Documentation UI + flow strip"),
            ("Jira scoped to [QA] workspace", "Shipped", "test_jira_qa_routing_integration.py"),
            ("Setup / catalog health resilience", "Shipped", "setup_status_safe.py, BackendHealthBanner"),
            ("Documentation product (business-only)", "Shipped", "test_doc_export.py, test_feature_quality.py"),
            ("Signoff pytest bundle", "50 tests green", "PLATFORM-ENGINEERING-SIGNOFF-100.md §E"),
            ("Full backend regression", "807+ tests", "backend pytest (developer venv)"),
            ("User launch copy", "Ready", "scripts/export-platform-copy.ps1 → Downloads + LAUNCH.md"),
            ("Live tenant QA sync", "Ops", "Setup → select Jetts/LIFT [QA] → Sync"),
            ("Staging cert / burn-in", "Ops", "PLATFORM-P0-P1-FIX-PLAN.md §4.2"),
        ],
    )

    add_heading(doc, "How Fabric Beacon fits your stack", 1)
    add_table(
        doc,
        ["Concept", "What it means here"],
        [
            ("Catalog-first UX", "Analysts browse and search without opening Power BI Service for every question."),
            ("Scoped Jira intelligence", "Ticket text → client [QA] workspace → matches only inside that scope."),
            ("Optional agentic path", "Cursor MCP (pbi-discovery, discovery-guide) — not required for web analysts."),
            ("Human-in-the-loop", "Secrets in root .env only; build/GitHub delivery dry-run by default; control-plane key in prod."),
            ("Guardrails", "Rate limits, encrypted platform config, no browser token forms on /jira."),
            ("Sibling products", "Paginate Forge (paginated RDL), Release Orchestrator, optional label automation (operators)."),
        ],
    )

    add_heading(doc, "New user journey", 1)
    if journey_png and journey_png.is_file():
        doc.add_picture(str(journey_png), width=Inches(6.8))
    doc.add_paragraph(
        "Always launch .\\start-dev.ps1 (or -Demo) from the repo or Downloads copy — not frontend-only — "
        "so /api/setup/status and health checks succeed."
    )
    add_table(
        doc,
        ["Step", "Screen", "URL"],
        [
            ("1 — Learn", "About → Get started", "http://localhost:5173/about"),
            ("2 — Catalog", "Setup → Connect, workspaces, Sync", "http://localhost:5173/setup"),
            ("3 — Ticket", "Jira Analyzer", "http://localhost:5173/jira"),
            ("4 — Explore", "Discovery / Documentation", "http://localhost:5173/ · /documentation"),
        ],
    )

    add_heading(doc, "Feature highlights", 1)
    add_table(
        doc,
        ["Area", "What you get"],
        [
            ("Discovery", "Workspace-scoped reports, paginated reports, dashboards (when tenant has them)."),
            ("Inspector", "Tables, columns, measures, DAX, relationships, RLS roles."),
            ("Search & reverse lookup", "Find fields and measures across the synced catalog."),
            ("Lineage", "Model-to-report graph where Scanner/metadata allows."),
            ("Governance", "Health scores, stale datasets, workspace-filtered stats."),
            ("Documentation", "Business narratives, Generate latest, collapsed quality panel."),
            ("Jira Analyzer", "[QA] scope only; quality badge; blocked state when scope missing."),
            ("Ops", "/ops dashboard + GET /api/platform/status for operators."),
        ],
    )

    add_heading(doc, "First-time setup (plain language)", 1)
    steps = [
        "Clone the repo or use the Downloads folder: %USERPROFILE%\\Downloads\\pbi-discovery-platform",
        "copy .env.example .env at the repository root (one file for all secrets).",
        "Demo (no Azure): .\\start-dev.ps1 -Demo",
        "Live catalog: fill PBI_PBI_* and optional PBI_JIRA_* in .env → .\\start-dev.ps1",
        "Open About → follow the four-step strip → complete Setup sync for your [QA] workspaces.",
        "Restart the backend after any .env change.",
    ]
    for i, text in enumerate(steps, start=1):
        doc.add_paragraph(f"{i}. {text}", style="List Number")

    add_heading(doc, "Verification & artifacts", 1)
    add_table(
        doc,
        ["Artifact", "Purpose"],
        [
            ("docs/DOCUMENTATION.md", "Single onboarding reference for all users"),
            ("docs/PLATFORM-ENGINEERING-SIGNOFF-100.md", "Formal engineering GO — 100% scope"),
            ("docs/RELEASE-COPY-SIGNOFF.md", "Checklist before tagging a release copy"),
            ("docs/PLATFORM-P0-P1-FIX-PLAN.md", "Remaining ops + catalog depth tracking"),
            ("LAUNCH.md (Downloads copy)", "First-run steps in the exported user folder"),
            ("docs/SLACK-ANNOUNCEMENT-DRAFT.md", "Paste-ready Slack text + links"),
        ],
    )

    add_heading(doc, "Slack & Confluence", 1)
    doc.add_paragraph(
        "Post the link to docs/DOCUMENTATION.md in your repo (not secrets). "
        "Use docs/SLACK-ANNOUNCEMENT-DRAFT.md for the full announcement. "
        "Attach this Word file for leadership summaries."
    )

    add_heading(doc, "Regenerate this document", 1)
    doc.add_paragraph(
        "From the repository root (after pip install python-docx matplotlib):\n"
        "python scripts/generate-fabric-beacon-insights-docx.py"
    )

    return doc


def main() -> int:
    for path in (OUT_DOCS, OUT_DOWNLOADS):
        if path.is_file():
            path.unlink()

    journey_png: Path | None = None
    try:
        ASSETS.mkdir(parents=True, exist_ok=True)
        journey_png = ASSETS / "user-journey.png"
        draw_user_journey(journey_png)
    except Exception as exc:
        print(f"Note: journey diagram skipped ({exc})", file=sys.stderr)
        journey_png = None

    doc = build_docx(journey_png)
    OUT_DOCS.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCS)
    doc.save(OUT_DOWNLOADS)
    print(f"Wrote {OUT_DOCS}")
    print(f"Wrote {OUT_DOWNLOADS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
